#!/usr/bin/env python3
"""Smoke test of pro-pptx-studio (MIT License): run it after changing the skill, before shipping a
new version. It builds one small French deck on a real template of the user's library and checks that
every step still works: build, polish, rule check, render.

  python3 tests/smoke.py --lib LIB [--template ID] [--out DIR]

It never runs while a presentation is being made, and it does not need a test template: it uses a
template from the library (the first 16:9 one not rated "no", or --template). Prints PASS/FAIL per
check and exits with 1 when something fails. Takes about one minute.
"""
import argparse, json, os, re, subprocess, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), 'scripts')
sys.path.insert(0, SCRIPTS)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--lib', default=os.environ.get('PPS_LIB'))
    ap.add_argument('--template')
    ap.add_argument('--out', default=tempfile.mkdtemp(prefix='pps-smoke-'))
    a = ap.parse_args()
    if not a.lib:
        sys.exit('give --lib or set PPS_LIB')
    lib = os.path.abspath(a.lib)
    idx = json.load(open(os.path.join(lib, 'index.json')))['templates']
    import tplbuild
    from tplpreview import cover_plan, strip_slides
    tid = a.template
    if not tid:
        for t in sorted(idx.values(), key=lambda t: t['id']):
            if t.get('rating') != 'no' and abs(t.get('aspect', 0) - 16 / 9) < 0.05 and t.get('slides_count', 0) >= 5:
                try:
                    tplbuild.template_path(lib, t['id'])
                    tid = t['id']
                    break
                except Exception:
                    continue
    if not tid:
        sys.exit('no usable 16:9 template found in the library')
    meta = json.load(open(os.path.join(lib, tid, 'meta.json')))
    plan, content, _ = strip_slides(meta)
    cn = content['n'] if content else 2
    last = meta['slides_count']
    text, remove = cover_plan(meta, {'title': 'Projet Test', 'subtitle': 'Vérification automatique du skill', 'meta': ['Équipe', '2026']}, lib, tid)
    os.makedirs(a.out, exist_ok=True)
    out = os.path.join(a.out, 'smoke.pptx')
    T = lambda t: {'kit': 'text', 'role': 'title', 'box': [0.08, 0.08, 0.84, 0.14], 'paras': t, 'name': 'smoke-title'}
    spec = {'lib': lib, 'base': tid, 'out': out, 'lang': 'fr', 'motion': 'auto', 'talk': {'minutes': 2},
            'slides': [
                {'use': '#1', 'text': text, 'remove': remove, 'notes': 'Bonjour à tous.'},
                {'use': '#%d' % cn, 'canvas': {'keep': []}, 'add': [T('Plan de la présentation'),
                 {'kit': 'bullets', 'box': [0.1, 0.3, 0.8, 0.55], 'items': ['Sprint 1 : analyse', 'Sprint 2 : tests à 5 GHz']}],
                 'notes': 'Voici le plan.'},
                {'use': '#%d' % cn, 'canvas': {'keep': []}, 'add': [T('Choix techniques'),
                 {'kit': 'code', 'box': [0.1, 0.3, 0.8, 0.55], 'lang': 'terminal', 'src': '$ terraform init\nInitializing...\n$ terraform apply'}],
                 'notes': 'Les commandes utilisées.'},
                {'use': '#%d' % cn, 'canvas': {'keep': []}, 'add': [T('Analyse des coûts'),
                 {'kit': 'chart', 'name': 'chart-cost', 'box': [0.15, 0.3, 0.7, 0.55], 'kind': 'column', 'categories': ['A', 'B'],
                  'series': [{'name': 'Coût', 'values': [10, 20]}], 'highlight': 1}], 'notes': 'Le coût.'},
                {'use': '#%d' % last, 'canvas': {'keep': []}, 'add': [T('Merci !'),
                 {'kit': 'qr', 'box': [0.42, 0.45, 0.16, 0.4], 'url': 'https://example.com/project', 'label': 'Dépôt'}],
                 'notes': 'Merci. Des questions ?'},
                {'use': '#%d' % cn, 'backup': True, 'canvas': {'keep': []}, 'add': [T('Question : sécurité ?')]},
            ],
            'polish': {'charts': [{'slide': 4, 'name': 'chart-cost', 'highlight': 1}]}}
    sp = os.path.join(a.out, 'smoke.json')
    json.dump(spec, open(sp, 'w'), indent=1, ensure_ascii=False)
    results = []

    def check(name, ok, detail=''):
        results.append(ok)
        print('%s  %s%s' % ('PASS' if ok else 'FAIL', name, (' - ' + detail) if detail and not ok else ''))

    env = dict(os.environ, PPS_LIB=lib)
    run = lambda *c: subprocess.run([sys.executable] + list(c), capture_output=True, text=True, env=env, timeout=900)
    r = run(os.path.join(SCRIPTS, 'tplbuild.py'), sp)
    check('build', r.returncode == 0 and os.path.exists(out), r.stderr[-400:])
    if r.returncode:
        sys.exit(1)
    r = run(os.path.join(SCRIPTS, 'tplpolish.py'), sp)
    check('polish', r.returncode == 0, r.stderr[-400:])
    from pptx import Presentation
    prs = Presentation(out)
    sl = list(prs.slides)
    check('slide count', len(sl) == 6, str(len(sl)))
    num = lambda s: next((x.text_frame.text.strip() for x in s.shapes if x.name.lstrip('!') in ('slide-number', 'slide-number-tpl') and x.has_text_frame), '')
    check('cover has no number', num(sl[0]) == '', num(sl[0]))
    check('plan is number 1', num(sl[1]).lstrip('0') == '1', num(sl[1]))
    check('closing is numbered', num(sl[4]).lstrip('0') == '4', num(sl[4]))
    check('backup slide is A1', num(sl[5]) == 'A1', num(sl[5]))
    alltext = '\n'.join(x.text_frame.text for s in sl for x in s.shapes if x.has_text_frame)
    check('word groups kept together', 'Sprint 1' in alltext and '5 GHz' in alltext)
    check('French spacing', 'Merci !' in alltext and 'Sprint 1 :' in alltext)
    cols = {r_.font.color.rgb for x in sl[2].shapes if x.name == 'code-1-text' or x.name.endswith('-text') and x.has_text_frame
            for p in x.text_frame.paragraphs for r_ in p.runs if r_.font.color and r_.font.color.type is not None}
    check('code is colored', len(cols) >= 3, str(len(cols)))
    check('Morph names', sum(1 for s in sl if any(x.name == '!!title' for x in s.shapes)) >= 5)
    check('QR code present', any(x.name.startswith('qr') and x.shape_type == 13 for x in sl[4].shapes))
    rd = os.path.join(a.out, 'render')
    r = run(os.path.join(SCRIPTS, 'render.py'), out, rd, '--dpi', '50')
    imgs = [f for f in os.listdir(rd) if re.match(r'slide-\d+\.jpg', f)] if os.path.isdir(rd) else []
    check('render', len(imgs) == 6, r.stderr[-300:])
    r = run(os.path.join(SCRIPTS, 'tplcheck.py'), sp, '--render', rd, '--json')
    try:
        rep = json.loads(r.stdout)
        check('rule check runs', True)
        bad = [i for i in rep['issues'] if i['rule'] in ('numbers', 'off-slide', 'sample-text')]
        check('no numbering / off-slide / sample-text issues', not bad, json.dumps(bad, ensure_ascii=False)[:400])
        check('talk time estimated', bool(rep['talk'].get('estimate')))
    except Exception as e:
        check('rule check runs', False, repr(e) + r.stderr[-300:])
    print('\n%d/%d passed  (template %s, files in %s)' % (sum(results), len(results), tid, a.out))
    sys.exit(0 if all(results) else 1)


if __name__ == '__main__':
    main()

#!/usr/bin/env python3
"""Build a deck on top of real templates (MIT License).

  python3 tplbuild.py spec.json            -> writes spec["out"], prints a JSON report

The spec picks slides from one or more library templates, fills their text (fitted to the
designer's boxes), swaps pictures, copies shapes between templates, recolors and re-fonts the
whole deck for the topic, draws kit components (charts, KPIs, code, tables, mockups, timelines,
cut-out objects with glow) in the template's fonts and colors, adds speaker notes, Morph
transitions and entrance animations, and embeds the fonts. See references/template-mode.md.
"""
import copy, io, json, os, re, shutil, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from PIL import Image  # noqa: E402
from lxml import etree  # noqa: E402
from tplkit import pkg as P, text as T, recolor as RC, fonts as F, motion as M  # noqa: E402
from tplkit.pkg import q  # noqa: E402

LOREM = re.compile(r'(?i)lorem ipsum|consectetur adipiscing|your sample|reallygreatsite|presentation title|@really')


def lib_path(spec):
    return os.path.abspath(spec.get('lib') or os.environ.get('PPS_LIB') or os.path.join(os.path.dirname(HERE), 'library'))


class Lib:
    def __init__(self, root):
        self.root = root
        self.pk, self.meta = {}, {}

    def path(self, tid):
        return template_path(self.root, tid)

    def pkg(self, tid):
        if tid not in self.pk:
            self.pk[tid] = P.Pkg(self.path(tid))
            T.add_font_dir(os.path.join(self.root, tid, 'fonts'))
        return self.pk[tid]

    def m(self, tid):
        if tid not in self.meta:
            self.meta[tid] = json.load(open(os.path.join(self.root, tid, 'meta.json')))
        return self.meta[tid]


def template_path(root, tid):
    """<lib>/<id>/template.pptx, or the original file in the templates folder (PPS_TEMPLATES or
    <lib>/../Templates) when the library keeps only the index."""
    p = os.path.join(root, tid, 'template.pptx')
    if os.path.exists(p):
        return p
    meta = json.load(open(os.path.join(root, tid, 'meta.json')))
    for d in [os.environ.get('PPS_TEMPLATES'), os.path.join(os.path.dirname(root), 'Templates'), os.path.join(root, '..', 'templates')]:
        if d and os.path.exists(os.path.join(d, meta.get('source_file', ''))):
            return os.path.join(d, meta['source_file'])
    raise FileNotFoundError('template file for %s not found: put %s in the library folder or set PPS_TEMPLATES' % (tid, meta.get('source_file')))


def parse_ref(ref, base):
    tid, _, n = ref.partition('#')
    return (tid or base), int(n)


def replace_image(pkg, part, key, path):
    x = pkg.xml(part)
    sh = P.find_shape(x, key)
    if sh is None:
        return 'shape not found: %s' % key
    blip = sh.find('.//' + q('a:blip'))
    if blip is None:
        return 'no picture in %s' % key
    im = Image.open(path)
    ext = 'png' if (im.mode in ('RGBA', 'LA', 'P')) else 'jpeg'
    buf = io.BytesIO()
    if ext == 'jpeg':
        im = im.convert('RGB')
        if max(im.size) > 2400:
            im.thumbnail((2400, 2400))
        im.save(buf, 'JPEG', quality=90)
    else:
        if max(im.size) > 2400:
            im.thumbnail((2400, 2400))
        im.save(buf, 'PNG', optimize=True)
    name = pkg.new_name('ppt/media/user{}.' + ext)
    pkg.add_part(name, buf.getvalue())
    pkg.ensure_default(ext, 'image/png' if ext == 'png' else 'image/jpeg')
    rid = pkg.add_rel(part, P.RT + 'image', name)
    blip.set('{%s}embed' % P.R, rid)
    # crop to fill the frame
    box = P.shape_box(sh)
    bf = blip.getparent()
    for e in bf.findall(q('a:srcRect')):
        bf.remove(e)
    if box and box[2] and box[3]:
        fa, ia = box[2] / box[3], im.width / im.height
        sr = etree.Element(q('a:srcRect'))
        if ia > fa:
            c = int((1 - fa / ia) / 2 * 100000)
            sr.set('l', str(c)); sr.set('r', str(c))
        elif ia < fa:
            c = int((1 - ia / fa) / 2 * 100000)
            sr.set('t', str(c)); sr.set('b', str(c))
        blip.addnext(sr)
    pkg.touch(part)
    return None


def faces_by_role(meta):
    acc = {}
    for s in meta['slides']:
        for t in s['texts']:
            d = 0 if t['kind'] in ('title', 'number', 'step') or t['size'] >= 40 else 1
            acc.setdefault(t['face'], [0, 0])[d] += t['chars']
    disp = {f: v[0] for f, v in acc.items() if v[0] >= v[1]}
    body = {f: v[1] for f, v in acc.items() if v[1] > v[0]}
    return disp, body


def theme_from(meta, pkg, parts, font_map, overrides):
    disp, body = faces_by_role(meta)
    dface = max(disp, key=disp.get) if disp else 'Inter'
    bface = max(body, key=body.get) if body else 'Inter'
    dface, bface = font_map.get(dface, dface), font_map.get(bface, bface)
    dark = meta.get('mode') != 'light'
    cols = []
    for p in parts:
        cols += re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', etree.tostring(pkg.xml(p)).decode('utf8'))
    from collections import Counter
    sat = Counter()
    for c in cols:
        h, s, v = RC.hex_to_hsv(c)
        if s > 0.45 and v > 0.45:
            sat[c.upper()] += 1
    ranked = [c for c, _ in sat.most_common(8)]
    pal = [c for c in meta.get('palette', []) if RC.hex_to_hsv(c)[1] > 0.35 and RC.hex_to_hsv(c)[2] > 0.35]
    ranked = pal + [c for c in ranked if c not in pal]
    accent = ranked[0] if ranked else ('2D8CFF')
    accent2 = next((c for c in ranked[1:] if abs(RC.hex_to_hsv(c)[0] - RC.hex_to_hsv(accent)[0]) > 0.04), None)
    if not accent2:
        h, s, v = RC.hex_to_hsv(accent)
        accent2 = RC.hsv_to_hex(h + 0.06, s * 0.8, min(1, v * 1.15))
    th = {'display': dface, 'body': bface, 'mono': 'JetBrains Mono', 'text': 'FFFFFF' if dark else '161A22',
          'muted': 'A9B4C8' if dark else '5B6476', 'accent': accent, 'accent2': accent2, 'surface': 'FFFFFF' if dark else '0E1A2B',
          'surface_alpha': 8 if dark else 5, 'border_alpha': 24 if dark else 14, 'dark': dark}
    th.update(overrides or {})
    return th


TITLE_FRACTION = 0.035   # text at least this share of the slide width (in pt) is treated as a headline


def build(spec):
    lib = Lib(lib_path(spec))
    base = spec['base']
    lang = spec.get('lang', 'en')
    # word groups on one line ("Sprint 1", "10 Aug", "5 GHz"); French spacing before ? ! ; : when lang is fr
    spec = dict(spec, replace={k: T.typography(v, lang) for k, v in spec.get('replace', {}).items()},
                slides=[dict(s_, text=T.typography_obj(s_.get('text', {}), lang),
                             add=[a_ if a_.get('kit') in ('code', 'qr') else T.typography_obj(a_, lang) for a_ in s_.get('add', [])])
                        for s_ in spec['slides']])
    out = os.path.abspath(spec['out'])
    tmpdir = tempfile.mkdtemp(prefix='tplb_')
    report = {'out': out, 'warnings': [], 'fits': []}
    dst = P.Pkg(lib.path(base))
    T.add_font_dir(os.path.join(lib.root, base, 'fonts'))
    T.add_font_dir(F.SKILL_FONTS)
    T.add_font_dir(os.path.join(lib.root, '_fonts'))
    base_parts = dst.slides()
    used_sources = {base}
    new_parts = []
    # ---------- phase A: structure ----------
    slide_src = []
    for i, s in enumerate(spec['slides']):
        tid, n = parse_ref(s['use'], base)
        used_sources.add(tid)
        slide_src.append(tid)
        if tid == base:
            part = dst.duplicate_slide(base_parts[n - 1])
        else:
            src = lib.pkg(tid)
            part = P.copy_slide(src, src.slides()[n - 1], dst)
        new_parts.append(part)
    dst.set_slides(new_parts)
    # ---------- phase A2: content ----------
    W_, H_ = dst.size()
    headlines = {}   # part -> [(shape key, slot, id, max_lines)] fitted after the font swap
    for i, (s, part) in enumerate(zip(spec['slides'], new_parts), 1):
        x = dst.xml(part)
        if s.get('canvas'):
            # keep the background, small decorations and the listed shapes; remove the rest
            keep = set(s['canvas'].get('keep', [])) if isinstance(s['canvas'], dict) else set()
            for sh in list(P.top_shapes(x)):
                nm = P.shape_name(sh)
                if nm in keep:
                    continue
                b = P.shape_box(sh)
                area = (b[2] * b[3]) / (W_ * H_) if b else 0
                has_text = bool(P.shape_text(sh).strip())
                is_frame = sh.tag.endswith('graphicFrame')
                if area >= 0.5 and not has_text and not is_frame:
                    continue          # background art
                if not has_text and not is_frame and area < 0.12:
                    # small decoration (corner marks, lines, glows) unless it sits where kit content goes
                    hit = False
                    for a in s.get('add', []):
                        bx = a.get('box')
                        if bx and b and all(0 <= v <= 1.0001 for v in bx):
                            ax, ay, aw, ah = bx[0] * W_, bx[1] * H_, bx[2] * W_, bx[3] * H_
                            if b[0] < ax + aw and b[0] + b[2] > ax and b[1] < ay + ah and b[1] + b[3] > ay:
                                hit = True
                    if not hit:
                        continue
                sh.getparent().remove(sh)
        for old_t, new_t in spec.get('replace', {}).items():
            for t in x.iter(q('a:t')):
                if t.text and t.text.strip() == old_t:
                    t.text = new_t
        for ci, c in enumerate(s.get('copy', []), 1):
            tid, n = parse_ref(c['from'], base)
            src = lib.pkg(tid) if tid != base else P.Pkg(lib.path(base))
            sp = src.slides()[n - 1]
            W, H = dst.size()
            SW, SH = src.size()
            box = c.get('box')
            if box and all(0 <= v <= 1.0001 for v in box):
                box = [box[0] * W, box[1] * H, box[2] * W, box[3] * H]
            prefix = c.get('prefix', 'c%d:' % ci)
            if c.get('region'):
                rg = c['region']
                if all(0 <= v <= 1.0001 for v in rg):
                    rg = [rg[0] * SW, rg[1] * SH, rg[2] * SW, rg[3] * SH]
                els = P.copy_region(src, sp, rg, dst, part, box or [0, 0, W, H], prefix=prefix)
            else:
                sh = P.find_shape(src.xml(sp), c['shape'])
                if sh is None:
                    report['warnings'].append('slide %d: copy: no shape %s in %s' % (i, c['shape'], c['from']))
                    continue
                el = P.copy_shape(src, sp, sh, dst, part, box=box)
                for cc in el.iter(q('p:cNvPr')):
                    cc.set('name', prefix + cc.get('name', ''))
                if c.get('name'):
                    el.find('.//' + q('p:cNvPr')).set('name', c['name'])
                els = [el]
            if c.get('adapt_text'):
                P.adapt_text(els, on=c['adapt_text'])
            report.setdefault('copied', []).append({'slide': i, 'from': c['from'], 'shapes': [P.shape_name(e) for e in els]})
            x = dst.xml(part)
        for key in s.get('remove', []):
            if not P.remove_shape(x, key):
                report['warnings'].append('slide %d: remove: no shape %s' % (i, key))
        for key, val in s.get('text', {}).items():
            ml, mins, fit, merged = None, 0.55, 'auto', False
            if isinstance(val, dict):
                ml = val.get('max_lines')
                mins = val.get('min_scale', 0.55)
                fit = val.get('fit', 'auto')
                if val.get('merge'):
                    # one headline over the area of several template title lines (e.g. a long title
                    # in a two-shape "Digital / Media" cover): grow this box over them, drop them
                    me = P.find_shape(x, key)
                    boxes = [P.shape_box(me)] if me is not None else []
                    for other in val['merge']:
                        o = P.find_shape(x, other)
                        if o is not None and P.shape_box(o):
                            boxes.append(P.shape_box(o))
                            o.getparent().remove(o)
                    if me is not None and len(boxes) > 1:
                        x0, y0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
                        x1, y1 = max(b[0] + b[2] for b in boxes), max(b[1] + b[3] for b in boxes)
                        P.set_box(me, [x0, y0, x1 - x0, y1 - y0])
                        fit = 'title'
                        ml = ml or 4
                        merged = True
                val = val['text']
            rep = T.set_text(x, key, val, max_lines=ml or s.get('max_lines', {}).get(key), min_scale=mins,
                             fit=fit, title_pt=W_ / 12700 * TITLE_FRACTION, trust_box=merged)
            rep['slide'] = i
            if rep.get('title'):
                headlines.setdefault(part, []).append((key, rep['slot'], rep['id'], ml))
                continue
            report['fits'].append(rep)
            if 'warning' in rep or 'error' in rep:
                report['warnings'].append('slide %d: %s: %s' % (i, key, rep.get('warning') or rep.get('error')))
        for key, path in s.get('images', {}).items():
            err = replace_image(dst, part, key, path)
            if err:
                report['warnings'].append('slide %d: %s' % (i, err))
        if s.get('clear_lorem', True):
            for sh in list(P.top_shapes(x)) + list(x.iter(q('p:sp'))):
                if sh.getparent() is None:
                    continue
                t = P.shape_text(sh)
                if t and LOREM.search(t) and P.shape_name(sh) not in s.get('text', {}):
                    try:
                        sh.getparent().remove(sh)
                    except Exception:
                        pass
        dst.touch(part)
    # ---------- recolor and fonts ----------
    groups = {}
    for tid, part in zip(slide_src, new_parts):
        groups.setdefault(tid, []).append(part)
    if spec.get('recolor'):
        rc = spec['recolor'] if isinstance(spec['recolor'], dict) else {'to': spec['recolor']}
        report['recolor'] = []
        for tid, parts in groups.items():
            sp_ = lib.pkg(tid) if tid != base else P.Pkg(lib.path(base))
            h = RC.accent_hue(sp_, sp_.slides())
            if h is None:
                h = RC.image_accent_hue(dst, parts)   # templates whose color lives in the pictures
            report['recolor'].append(dict(RC.recolor(dst, parts, rc['to'], images=rc.get('images', 'auto'), sat_mul=rc.get('sat', 1.0), src_hue=h, theme=(tid == base)), template=tid))
    elif spec.get('harmonize', True) and len(groups) > 1:
        bp = P.Pkg(lib.path(base))
        bh = RC.accent_hue(bp, bp.slides())
        if bh is not None:
            target = RC.hsv_to_hex(bh, 0.85, 0.9)
            report['harmonize'] = []
            for tid, parts in groups.items():
                if tid == base:
                    continue
                sp = lib.pkg(tid)
                h = RC.accent_hue(sp, sp.slides())
                hi = RC.image_accent_hue(dst, parts)
                if h is None or (hi is not None and abs(((hi - bh) + 0.5) % 1 - 0.5) > 0.08):
                    h = hi if h is None else h
                    report.setdefault('harmonize_images', []).append(tid)
                if h is not None and (abs(((h - bh) + 0.5) % 1 - 0.5) > 0.03 or tid in report.get('harmonize_images', [])):
                    report['harmonize'].append(dict(RC.recolor(dst, parts, target, images='auto', src_hue=hi if tid in report.get('harmonize_images', []) else h, theme=False, mode='mono' if tid in report.get('harmonize_images', []) else 'rotate'), template=tid))
    for s_, part in zip(spec['slides'], new_parts):
        if s_.get('recolor'):
            rc = s_['recolor'] if isinstance(s_['recolor'], dict) else {'to': s_['recolor']}
            RC.recolor(dst, [part], rc['to'], images=rc.get('images', 'auto'), mode=rc.get('mode', 'rotate'), theme=False)
    font_map = dict(spec.get('fonts', {}).get('map', {}))
    role_fonts = None
    if spec.get('fonts', {}).get('theme') or spec.get('fonts', {}).get('display'):
        if spec['fonts'].get('theme'):
            key, _, k = spec['fonts']['theme'].partition(':')
            role_fonts = F.THEME_FONTS[key][int(k or 0)]
        else:
            role_fonts = (spec['fonts']['display'], spec['fonts'].get('body', spec['fonts']['display']))
    if font_map or role_fonts:
        bold_map = {k: True for k in font_map if re.search(r'(?i)bold|black|heavy', k)}
        before = {part: T.footprints(dst.xml(part)) for part in new_parts}
        if font_map:
            F.swap_fonts(dst, new_parts, font_map, bold_map)
        if role_fonts:
            F.swap_by_role(dst, new_parts, role_fonts[0], role_fonts[1], keep=set(spec['fonts'].get('keep', [])))
            font_map = {'display': role_fonts[0], 'body': role_fonts[1], **font_map}
        for i, part in enumerate(new_parts, 1):
            ch = T.keep_footprints(dst.xml(part), before[part], skip={t_[2] for t_ in headlines.get(part, [])})
            if ch:
                report.setdefault('refit_after_font_swap', []).append({'slide': i, 'shapes': ch})
                for c_ in ch:
                    if len(c_) > 2:
                        report['warnings'].append('slide %d: %s: %s (the new font is wider)' % (i, c_[0], c_[2]))
            dst.touch(part)
    report['font_map'] = font_map
    # headlines: balanced lines, sized to fill their slot (runs after the font swap, with the final font)
    for i, part in enumerate(new_parts, 1):
        # lowest first, so a title above sees where the one below really starts
        for key, slot, sid, ml in sorted(headlines.get(part, []), key=lambda h: -h[1]['top']):
            rep = T.fit_title(dst.xml(part), '#' + sid, slot, (W_, H_), max_lines=ml or 4)
            rep.update(slide=i, shape=key, id='#' + sid)
            report['fits'].append(rep)
            if 'warning' in rep:
                report['warnings'].append('slide %d: %s: %s' % (i, key, rep['warning']))
            if rep.get('overlaps'):
                report['warnings'].append('slide %d: %s now runs under %s (placed over the sample title): check it, move or remove' % (i, key, ', '.join(repr(n) for n in rep['overlaps'])))
        # title lines split over shapes of one template size get one size (the smallest fit)
        fitted = [f for f in report['fits'] if f.get('slide') == i and f.get('size_pt')]
        for f in fitted:
            same = [g for g in fitted if abs(g['template_pt'] - f['template_pt']) < 0.01 * f['template_pt']]
            m = min(g['size_pt'] for g in same)
            if len(same) > 1 and f['size_pt'] > m:
                T.set_size(dst.xml(part), f['id'], m)
                f['size_pt'] = m
        if headlines.get(part):
            dst.touch(part)
    # one line per slide for tplpolish (Morph names) and tplcheck (rules): template role, title shape, backup flag
    info = []
    for i, (s, tid) in enumerate(zip(spec['slides'], slide_src), 1):
        _, n = parse_ref(s['use'], base)
        ms = lib.m(tid)['slides'][n - 1]
        heads = [f for f in report['fits'] if f.get('slide') == i and f.get('template_pt') and f.get('id')]
        title = None
        if s.get('title_shapes'):
            sh = P.find_shape(dst.xml(new_parts[i - 1]), s['title_shapes'][0])
            title = '#' + P.shape_id(sh) if sh is not None else s['title_shapes'][0]   # a kit shape: by name
        elif heads:
            title = max(heads, key=lambda f: f['template_pt'])['id']
        info.append({'slide': i, 'use': '%s#%d' % (tid, n), 'role': s.get('role') or ms.get('role_override') or ms['role'],
                     'title': title, 'backup': bool(s.get('backup'))})
    report['slides_info'] = info
    report['lang'] = lang
    th = theme_from(lib.m(base), dst, new_parts, {} if role_fonts else font_map, spec.get('theme'))
    if role_fonts:
        th['display'], th['body'] = role_fonts
    report['theme'] = th
    mid = os.path.join(tmpdir, 'stage_a.pptx')
    dst.save(mid)
    # ---------- phase B: kit components and notes (python-pptx) ----------
    from pptx import Presentation
    from tplkit import kit as K
    prs = Presentation(mid)
    K.SCALE = prs.slide_width / 12192000
    theme = K.Theme.from_dict(th)
    credit_needed = [t for t in used_sources if lib.m(t).get('source') == 'slidescarnival']
    for i, (s, slide) in enumerate(zip(spec['slides'], prs.slides), 1):
        for j, a in enumerate(s.get('add', [])):
            a = dict(a)
            kind = a.pop('kit')
            nm = a.pop('name', '%s-%d' % (kind, j + 1))
            box = a.pop('box')
            fn = getattr(K, kind)
            if kind in ('image', 'depth', 'mockup'):
                a.setdefault('tmpdir', tmpdir)
            if kind == 'text' and a.get('role') == 'title' and not report['slides_info'][i - 1]['title']:
                report['slides_info'][i - 1]['title'] = nm      # a title drawn with the kit (canvas slides)
            try:
                fn(prs, slide, theme, box, name=nm, **a)
            except Exception as e:
                report['warnings'].append('slide %d: kit %s failed: %r' % (i, kind, e))
        if s.get('notes'):
            slide.notes_slide.notes_text_frame.text = s['notes']
        if credit_needed and i == len(spec['slides']) and spec.get('credit', True):
            K.credit(prs, slide, theme, 'Template: SlidesCarnival (slidescarnival.com), CC BY 4.0')
    prs.save(mid)
    # ---------- phase C: motion and fonts ----------
    dst = P.Pkg(mid)
    parts = dst.slides()
    mode = spec.get('motion', 'auto')
    W, H = dst.size()
    for i, (s, part) in enumerate(zip(spec['slides'], parts), 1):
        x = dst.xml(part)
        anim = s.get('animate', 'auto' if mode != 'none' else 'none')
        steps = []
        if anim == 'auto':
            titles = set(s.get('title_shapes', [])) | {k for k, r in ((r['shape'], r) for r in report['fits'] if r.get('slide') == i) if k in s.get('text', {}) and list(s['text']).index(k) == 0}
            steps = M.auto_steps(x, title_names=titles, slide_size=(W, H))[:12] if i > 1 else []
        elif isinstance(anim, list):
            steps = anim
        miss = M.apply(x, steps, transition=s.get('transition', spec.get('transition', 'morph')), mode='click' if mode == 'click' else 'auto')
        if miss:
            report['warnings'].append('slide %d: animate: no shapes %s' % (i, miss))
        dst.touch(part)
    report['embedded_fonts'] = F.embed_skill_fonts(dst)
    dst.save(out)
    report['slides'] = len(parts)
    shutil.rmtree(tmpdir, ignore_errors=True)
    # leftovers
    for i, part in enumerate(P.Pkg(out).slides(), 1):
        pass
    return report


if __name__ == '__main__':
    spec = json.load(open(sys.argv[1]))
    rep = build(spec)
    rp = os.path.splitext(rep['out'])[0] + '.report.json'
    json.dump(rep, open(rp, 'w'), indent=1, ensure_ascii=False)
    print(json.dumps({'out': rep['out'], 'slides': rep.get('slides'), 'warnings': rep['warnings'],
                      'refits': rep.get('refit_after_font_swap', []), 'embedded_fonts': rep.get('embedded_fonts'),
                      'full_report': rp}, indent=1, ensure_ascii=False))

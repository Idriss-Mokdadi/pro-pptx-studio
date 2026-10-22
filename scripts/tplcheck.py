#!/usr/bin/env python3
"""tplcheck.py - checks a built deck against the house rules and estimates the talk time (MIT License).

  python3 tplcheck.py spec.json [--render DIR] [--minutes 10] [--wpm 120] [--json]
  python3 tplcheck.py deck.pptx [...]

Run it after tplbuild.py + tplpolish.py (and after render.py when --render is given). It reads the
deck, its .report.json (title shape, template role and backup flag of every slide) and the spec, then
lists every broken rule with a level:
  fix    break a house rule: correct it and rebuild
  check  probably wrong: look at the slide and decide
  info   to know (values still to fill, talk-time plan)
Rules: text size, title at the top, bullets not paragraphs, at most 5 bullets, leftover sample text,
text off the slide, text over a picture, slide numbers (cover blank, numbered from the plan, backup
slides A1...), speaker notes on every slide, the same template slide twice in a row, one background
family (needs --render), and the talk time from the notes (spoken words / words per minute).
"""
import json, os, re, sys
from pptx import Presentation
from pptx.oxml.ns import qn

SAMPLE = re.compile(r'(?i)lorem ipsum|consectetur adipiscing|reallygreatsite|@really|borcelle|larana|your (company|name|text|title|sample)|'
                    r'presentation title|company name|presenter name|add a (main point|title|subtitle)|click to (add|edit)|insert (your|text)|'
                    r'write (your|a) |123-456-7890|hello@|www\.\w+site|elaborate on|briefly (describe|elaborate)')
TODO = re.compile(r'\[(value|to fill|tbd|todo|à compléter)\]', re.I)
SENT = re.compile(r'[.!?]\s+[A-ZÀ-Ý0-9]')
GROUP = re.compile(r'(?i)\b(sprint|phase|step|week|day|étape|semaine|jour)[ ]\d+\b|\b\d+[ ](%|GHz|MHz|ms|min|GB|MB|km|kW|dB)\b')
NO_TITLE_RULE = {'cover', 'closing', 'section', 'statement', 'quote', 'big-number', 'image', 'video', 'credits'}


def base(name):
    return name[2:] if name.startswith('!!') else name


def words(t):
    return len(re.findall(r"[\w'’\-]+", t))


def run_sizes(shape):
    """Font sizes (pt) of the runs in a text shape, with PowerPoint's shrink-on-overflow applied."""
    body = shape._element.find('.//' + qn('p:txBody'))
    if body is None:
        return []
    scale = 1.0
    na = body.find('.//' + qn('a:normAutofit'))
    if na is not None and na.get('fontScale'):
        scale = int(na.get('fontScale')) / 100000
    out = []
    for p in body.findall(qn('a:p')):
        dflt = p.find(qn('a:pPr') + '/' + qn('a:defRPr'))
        for r in p.findall(qn('a:r')):
            t = r.find(qn('a:t'))
            if t is None or not (t.text or '').strip():
                continue
            rpr = r.find(qn('a:rPr'))
            sz = rpr.get('sz') if rpr is not None else None
            sz = sz or (dflt.get('sz') if dflt is not None else None)
            if sz:
                out.append((int(sz) / 100 * scale, t.text))
    return out


def box(s, W, H):
    if s.left is None or s.width is None:
        return None
    return (s.left / W, s.top / H, s.width / W, s.height / H)


def overlap(a, b):
    return a[0] < b[0] + b[2] and a[0] + a[2] > b[0] and a[1] < b[1] + b[3] and a[1] + a[3] > b[1]


def main():
    a = sys.argv[1:]
    src = a[0]
    opt = lambda k, d=None: a[a.index(k) + 1] if k in a else d
    spec = {}
    if src.endswith('.json'):
        spec = json.load(open(src))
        deck = spec['out']
    else:
        deck = src
    rp = os.path.splitext(deck)[0] + '.report.json'
    report = json.load(open(rp)) if os.path.exists(rp) else {}
    info = {x['slide']: x for x in report.get('slides_info', [])}
    talk = dict(spec.get('talk', {}))
    if opt('--minutes'):
        talk['minutes'] = float(opt('--minutes'))
    wpm = float(opt('--wpm', talk.get('wpm', 120)))
    prs = Presentation(deck)
    W, H = prs.slide_width, prs.slide_height
    k = W / 12700 / 1440                       # sizes in the rules are for a 1440 pt-wide canvas
    min_read, min_hard = 16 * k, 12 * k
    slides = list(prs.slides)
    backup = {i for i, x in info.items() if x.get('backup')} | {i for i, s in enumerate(spec.get('slides', []), 1) if s.get('backup')}
    main_slides = [i for i in range(1, len(slides) + 1) if i not in backup]
    closing = main_slides[-1] if main_slides else len(slides)
    numbers = (spec.get('polish', {}) or {}).get('numbers', {})
    first_num = numbers.get('start_slide', 2) if isinstance(numbers, dict) else None
    issues = []

    def add(n, level, rule, msg):
        issues.append({'slide': n, 'level': level, 'rule': rule, 'message': msg})

    has_numbers = any(base(s.name) in ('slide-number', 'slide-number-tpl') for sl in slides for s in sl.shapes)
    prev_use = None
    for n, sl in enumerate(slides, 1):
        inf = info.get(n, {})
        role = inf.get('role') or ('cover' if n == 1 else 'closing' if n == closing else None)
        tref = inf.get('title') or ''
        shapes = list(sl.shapes)
        title = next((s for s in shapes if base(s.name) == 'title' or (tref.startswith('#') and str(s.shape_id) == tref[1:])
                      or (tref and not tref.startswith('#') and base(s.name) == tref)), None)
        # title at the top
        if title is not None and n != 1 and n != closing and role not in NO_TITLE_RULE and n not in backup:
            b = box(title, W, H)
            if b and b[1] > 0.30:
                add(n, 'fix', 'title-top', 'title starts at %d%% of the slide height: move it to the top or use another template slide' % round(b[1] * 100))
        pics = [s for s in shapes if s.shape_type == 13 and box(s, W, H) and not base(s.name).startswith(('depth', 'logo', 'qr', 'image', 'mockup', 'icon', 'node-'))]
        for s in shapes:
            nm = base(s.name)
            if not s.has_text_frame or nm.startswith(('slide-number', 'credit', 'code-', 'qr-')) or '-file' in nm:
                continue
            txt = s.text_frame.text
            if not txt.strip():
                continue
            b = box(s, W, H)
            # sample text and values to fill
            if SAMPLE.search(txt):
                add(n, 'fix', 'sample-text', 'template sample text left in %r: "%s"' % (s.name, txt.strip()[:50]))
            if TODO.search(txt):
                add(n, 'info', 'to-fill', 'value to fill in %r: "%s"' % (s.name, txt.strip()[:50]))
            # sizes
            small = [sz for sz, t in run_sizes(s) if sz < min_read]
            if small:
                lvl = 'fix' if min(small) < min_hard else 'check'
                add(n, lvl, 'text-size', '%r has text at %.0f pt (minimum %.0f pt on this slide size): enlarge it or shorten the text' % (s.name, min(small), min_read))
            # bullets, not paragraphs
            if s is not title and not nm.startswith(('kpi', 'table', 'chart')):
                paras = [p for p in txt.split('\n') if p.strip()]
                if len(paras) > 5 and not nm.startswith('t:'):
                    add(n, 'check', 'bullets', '%r has %d lines of text: keep at most 5 bullets, move the rest to the notes' % (s.name, len(paras)))
                for p in paras:
                    w = words(p)
                    if w >= 25 or (w >= 14 and SENT.search(p)):
                        add(n, 'fix', 'paragraph', '%r holds a paragraph (%d words): cut it to a short phrase, put the explanation in the notes' % (s.name, w))
                        break
                    if w > 12:
                        add(n, 'check', 'long-bullet', '%r: "%s..." has %d words (aim for 8 or fewer)' % (s.name, p[:40], w))
                        break
            if GROUP.search(txt):
                add(n, 'check', 'word-group', '%r: "%s" can break across two lines: use a non-breaking space' % (s.name, GROUP.search(txt).group(0)))
            # off the slide
            if b and (b[0] < -0.01 or b[1] < -0.01 or b[0] + b[2] > 1.01 or b[1] + b[3] > 1.01):
                add(n, 'fix', 'off-slide', '%r runs off the slide' % s.name)
            # text over the central part of a picture
            if b:
                for p in pics:
                    pb = box(p, W, H)
                    if not (0.06 <= pb[2] * pb[3] <= 0.85):
                        continue
                    core = (pb[0] + pb[2] * 0.2, pb[1] + pb[3] * 0.2, pb[2] * 0.6, pb[3] * 0.6)
                    if overlap(b, core):
                        add(n, 'check', 'text-over-art', '%r sits over the picture %r: move the text beside it' % (s.name, p.name))
                        break
        # slide numbers
        if has_numbers and first_num:
            vals = [s.text_frame.text.strip() for s in shapes if base(s.name) in ('slide-number', 'slide-number-tpl') and s.has_text_frame]
            shown = [v for v in vals if v]
            skip = set(numbers.get('skip', [])) if isinstance(numbers, dict) else set()
            if n < first_num and shown:
                add(n, 'fix', 'numbers', 'this slide shows a number (%s): the cover stays blank' % shown[0])
            elif n >= first_num and n not in skip and not shown:
                add(n, 'fix', 'numbers', 'slide number missing (numbering starts at the plan)')
            elif n in backup and shown and not re.match(r'^[A-Z]\d+$', shown[0]):
                add(n, 'check', 'numbers', 'backup slide numbered %s: expected A1, A2...' % shown[0])
        elif not has_numbers and n == 2:
            add(n, 'fix', 'numbers', 'the deck has no slide numbers: run tplpolish.py')
        # notes
        notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ''
        if not notes.strip() and n not in backup:
            add(n, 'check', 'notes', 'no speaker notes: add what the presenter says')
        # same template slide twice in a row
        use = inf.get('use')
        if use and use == prev_use:
            add(n, 'check', 'repeat', 'same template slide as the previous one (%s): use another layout' % use)
        prev_use = use
    # one background family (needs the render)
    rdir = opt('--render')
    if rdir:
        from PIL import Image
        import numpy as np
        lum = {}
        for n in range(1, len(slides) + 1):
            f = os.path.join(rdir, 'slide-%d.jpg' % n)
            f = f if os.path.exists(f) else os.path.join(rdir, 'slide-%02d.jpg' % n)
            if os.path.exists(f):
                lum[n] = float(np.asarray(Image.open(f).convert('L').resize((64, 36))).mean() / 255)
        if lum:
            dark = sum(1 for v in lum.values() if v < 0.45) >= len(lum) / 2
            for n, v in lum.items():
                if (v < 0.45) != dark:
                    add(n, 'fix', 'background', 'this slide is %s while the deck is %s: keep one background family' % ('dark' if v < 0.45 else 'light', 'dark' if dark else 'light'))
    # talk time from the notes (backup slides are not presented)
    per, cum, plan = [], 0.0, []
    for n, sl in enumerate(slides, 1):
        if n in backup:
            continue
        w = words(sl.notes_slide.notes_text_frame.text) if sl.has_notes_slide else 0
        sec = w / wpm * 60
        per.append((n, w, sec))
        plan.append({'slide': n, 'starts_at': '%d:%02d' % divmod(round(cum), 60)})
        cum += sec
        if sec > 110:
            add(n, 'check', 'talk-slide', 'notes take about %d:%02d on one slide: shorten them (do not split or merge slides unless the user asks)' % divmod(round(sec), 60))
    total_w = sum(w for _, w, _ in per)
    tk = {'words': total_w, 'wpm': wpm, 'estimate': '%d:%02d' % divmod(round(cum), 60)}
    if talk.get('minutes'):
        lim = float(talk['minutes']) * 60
        tk['allowed'] = '%d:%02d' % divmod(round(lim), 60)
        if cum > lim * 1.05:
            extra = round((cum - lim * 0.97) / 60 * wpm)
            add(0, 'fix', 'talk-time', 'the notes take about %s for %s allowed: remove about %d words from the notes (shorten, do not merge slides unless the user asks)' % (tk['estimate'], tk['allowed'], extra))
        elif cum < lim * 0.7:
            add(0, 'info', 'talk-time', 'the notes take about %s for %s allowed: there is room for more explanation' % (tk['estimate'], tk['allowed']))
    tk['plan'] = plan
    order = {'fix': 0, 'check': 1, 'info': 2}
    issues.sort(key=lambda x: (order[x['level']], x['slide']))
    counts = {l: sum(1 for i in issues if i['level'] == l) for l in order}
    if '--json' in a:
        print(json.dumps({'deck': deck, 'slides': len(slides), 'counts': counts, 'issues': issues, 'talk': tk}, indent=1, ensure_ascii=False))
        return
    print('%s: %d slides, %d fix, %d check, %d info' % (os.path.basename(deck), len(slides), counts['fix'], counts['check'], counts['info']))
    for i in issues:
        print('  [%s] %s %s: %s' % (i['level'], ('slide %d' % i['slide']) if i['slide'] else 'deck', i['rule'], i['message']))
    print('talk: %s words at %d wpm = about %s%s' % (total_w, wpm, tk['estimate'], (' (allowed %s)' % tk['allowed']) if 'allowed' in tk else ''))
    print('timing: ' + ', '.join('slide %d at %s' % (p['slide'], p['starts_at']) for p in plan))


if __name__ == '__main__':
    main()

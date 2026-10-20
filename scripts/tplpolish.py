#!/usr/bin/env python3
"""tplpolish.py - finishing pass after tplbuild.py (MIT). Usage: python3 tplpolish.py spec.json
Reads spec["out"], its .report.json theme, and spec["polish"]; boxes are slide fractions [x, y, w, h]."""
import copy, json, re, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

if sys.argv[1] == '--icons':      # python3 tplpolish.py --icons screenshot.png '{"vm":[x,y,r],...}' outdir  (x, y, r in pixels)
    import numpy as np, os
    im = Image.open(sys.argv[2]).convert('RGB'); os.makedirs(sys.argv[4], exist_ok=True)
    for name, (x, y, r) in json.loads(sys.argv[3]).items():
        c = im.crop((x - r, y - r, x + r, y + r)); a = np.asarray(c).astype(int)
        ys, xs = np.where(((a.max(2) - a.min(2)) > 45) | (a.mean(2) < 140))
        ic = np.asarray(c.crop((xs.min() - 6, ys.min() - 6, xs.max() + 7, ys.max() + 7)))
        rgba = np.dstack([ic, np.where(ic.min(2) > 238, 0, 255).astype(np.uint8)])
        Image.fromarray(rgba).save(os.path.join(sys.argv[4], name + '.png'))
    print('icons saved to', sys.argv[4]); sys.exit()
spec = json.load(open(sys.argv[1]))
P = spec.get('polish', {})
out = spec['out']
th = json.load(open(out.rsplit('.', 1)[0] + '.report.json')).get('theme', {})
TH = {'display': th.get('display', 'Inter'), 'body': th.get('body', 'Inter'), 'mono': th.get('mono', 'JetBrains Mono'),
      'text': th.get('text', 'FFFFFF'), 'muted': th.get('muted', 'A9B4C8'),
      'accent': th.get('accent', '115BB7'), 'accent2': th.get('accent2', '0BBDDC'), 'dark': th.get('dark', True)}
TH.update(P.get('theme', {}))
prs = Presentation(out)
W, H = prs.slide_width, prs.slide_height
AL = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}


def X(f): return Emu(int(f * W))
def Y(f): return Emu(int(f * H))
def S(n): return prs.slides[n - 1]
def find(sl, name): return next(s for s in sl.shapes if s.name == name)


def no_autofit(tf):
    bp = tf._txBody.find(qn('a:bodyPr'))
    for c in list(bp):
        if c.tag in (qn('a:normAutofit'), qn('a:spAutoFit')):
            bp.remove(c)


def fill(sh, c, alpha=None, c2=None, angle=90):
    """Solid fill, or a two-stop gradient when c2 is given."""
    if c2:
        sh.fill.gradient(); sh.fill.gradient_angle = angle
        st = sh.fill.gradient_stops
        st[0].position, st[1].position = 0.0, 1.0
        st[0].color.rgb, st[1].color.rgb = RGBColor.from_string(c), RGBColor.from_string(c2)
        return
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(c)
    if alpha is not None:
        clr = sh.fill._xPr.find('.//' + qn('a:srgbClr'))
        etree.SubElement(clr, qn('a:alpha')).set('val', str(int(alpha * 1000)))


def line(sh, c=None, w=1.0, dash=None):
    if c is None:
        sh.line.fill.background(); return
    sh.line.color.rgb = RGBColor.from_string(c); sh.line.width = Pt(w)
    if dash:
        etree.SubElement(sh.line._get_or_add_ln(), qn('a:prstDash')).set('val', dash)


def text(sl, box, paras, size=18, color=None, font=None, bold=False, align='l', anchor='m', name=None, spacing=None):
    """paras: str | [str | {t, size, color, font, bold}]"""
    t = sl.shapes.add_textbox(X(box[0]), Y(box[1]), X(box[2]), Y(box[3]))
    tf = t.text_frame; tf.word_wrap = True
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Emu(0))
    for i, p in enumerate([paras] if isinstance(paras, str) else paras):
        p = {'t': p} if isinstance(p, str) else p
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = AL[p.get('align', align)]
        if spacing:
            par.space_after = Pt(spacing)
        r = par.add_run(); r.text = p['t']
        r.font.name = TH.get(p.get('font', font or 'body'), p.get('font', font or TH['body']))
        r.font.size = Pt(p.get('size', size)); r.font.bold = p.get('bold', bold)
        r.font.color.rgb = RGBColor.from_string(TH.get(p.get('color', color or 'text'), p.get('color', color or TH['text'])))
    if name:
        t.name = name
    return t


def chip(sl, box, files, name='logos'):
    """Light rounded card holding logos side by side (logos with dark text or white backgrounds stay readable)."""
    x, y, w, h = box
    c = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x), Y(y), X(w), Y(h))
    c.adjustments[0] = 0.25; fill(c, 'F4F7FB'); line(c); c.shadow.inherit = False; c.name = name
    pad = h * 0.18; lh = h - 2 * pad
    ws = [lh * H / W * (lambda s: s[0] / s[1])(Image.open(f).size) for f in files]
    gap = (w - sum(ws)) / (len(files) + 1); cx = x + gap
    for f, lw in zip(files, ws):
        sl.shapes.add_picture(f, X(cx), Y(y + pad), X(lw), Y(lh)).name = name + '-logo'
        cx += lw + gap


def center(sh, box, size=None):
    """Put a text shape exactly over its container and center it both ways."""
    sh.left, sh.top, sh.width, sh.height = X(box[0]), Y(box[1]), X(box[2]), Y(box[3])
    tf = sh.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True; no_autofit(tf)
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Emu(0))
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            if size:
                r.font.size = Pt(size)


def title_lines(sh, lines, size, box=None, spacing=95):
    """Rewrite a template title as explicit lines (keeps its run style), e.g. to keep it clear of background art."""
    if box:
        sh.left, sh.top, sh.width, sh.height = X(box[0]), Y(box[1]), X(box[2]), Y(box[3])
    tf = sh.text_frame; no_autofit(tf)
    p0 = tf.paragraphs[0]._p
    rpr = tf.paragraphs[0].runs[0]._r.find(qn('a:rPr'))
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0):
        if r.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
            p0.remove(r)
    ppr = p0.find(qn('a:pPr'))
    if ppr is None:
        ppr = etree.Element(qn('a:pPr')); p0.insert(0, ppr)
    for o in ppr.findall(qn('a:lnSpc')):
        ppr.remove(o)
    ln = etree.Element(qn('a:lnSpc')); etree.SubElement(ln, qn('a:spcPct')).set('val', str(spacing * 1000)); ppr.insert(0, ln)
    end = p0.find(qn('a:endParaRPr'))
    for i, t in enumerate(lines):
        if i:
            br = etree.Element(qn('a:br'))
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            (end.addprevious(br) if end is not None else p0.append(br))
        r = etree.Element(qn('a:r'))
        if rpr is not None:
            nr = copy.deepcopy(rpr); nr.set('sz', str(int(size * 100))); r.append(nr)
        etree.SubElement(r, qn('a:t')).text = t
        (end.addprevious(r) if end is not None else p0.append(r))


def cards(sl, c):
    """Numbered hexagon cards in a grid: items [[num, head, sub], ...]."""
    x0, y0, w0, h0 = c.get('box', [0.09, 0.36, 0.82, 0.52])
    items, cols = c['items'], c.get('cols', 2)
    rows = -(-len(items) // cols); gx, gy = 0.02, 0.035
    cw, ch = (w0 - gx * (cols - 1)) / cols, (h0 - gy * (rows - 1)) / rows
    hs = min(ch * 0.66, 0.095); hw = hs * H / W * 1.1
    for i, (num, head, sub) in enumerate(items):
        x, y = x0 + (i // rows) * (cw + gx), y0 + (i % rows) * (ch + gy)
        card = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x), Y(y), X(cw), Y(ch))
        card.adjustments[0] = 0.18; fill(card, c.get('card', '0E2A4F'), c2=c.get('card2', '07111F'), angle=0)
        line(card, c.get('border', '1F4F8A'), 1); card.shadow.inherit = False; card.name = 'card-%s' % num
        hx = sl.shapes.add_shape(MSO_SHAPE.HEXAGON, X(x + 0.018), Y(y + (ch - hs) / 2), X(hw), Y(hs))
        fill(hx, TH['accent2'], c2=TH['accent']); line(hx); hx.shadow.inherit = False
        text(sl, [x + 0.018, y + (ch - hs) / 2, hw, hs], num, c.get('num_size', 26), 'FFFFFF', 'display', True, 'c')
        tx = x + 0.036 + hw
        text(sl, [tx, y + ch * 0.15, cw - (tx - x) - 0.01, ch * 0.38], head, c.get('head_size', 22), 'FFFFFF', 'body', True)
        text(sl, [tx, y + ch * 0.53, cw - (tx - x) - 0.01, ch * 0.32], sub, c.get('sub_size', 18), c.get('sub_color', '93C1E4'))


def diagram(sl, d):
    """Nodes: tile (light card: icon + name + kind, resource-visualizer style) or panel (translucent box with text lines).
    Links: from/to = [x, y] or "id.l|r|t|b", optional via [[x, y]...], arrows both|end|start|none, dash."""
    nodes = {}
    for n in d.get('nodes', []):
        x, y, w, h = n['box']; nodes[n['id']] = n['box']
        kind = n.get('type', 'tile')
        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x), Y(y), X(w), Y(h))
        sh.adjustments[0] = n.get('radius', 0.08); sh.shadow.inherit = False; sh.name = 'node-' + n['id']
        if kind == 'tile':
            fill(sh, n.get('fill', 'F4F7FB'))
            line(sh, TH['accent2'] if n.get('accent') else None, 3)
            if n.get('icon'):
                iw, ih = Image.open(n['icon']).size
                mh, mw = h * 0.46, w * 0.55
                ph = mh; pw = ph * iw / ih * H / W
                if pw > mw:
                    pw = mw; ph = pw * W / H * ih / iw
                sl.shapes.add_picture(n['icon'], X(x + (w - pw) / 2), Y(y + h * 0.1 + (mh - ph) / 2), X(pw), Y(ph))
            text(sl, [x + 0.004, y + h * 0.62, w - 0.008, h * 0.17], n.get('name', ''), n.get('name_size', 13), '0F1B2D', 'body', True, 'c')
            text(sl, [x + 0.004, y + h * 0.80, w - 0.008, h * 0.14], n.get('kind', ''), n.get('kind_size', 12), '5E6A7A', 'body', False, 'c')
        else:
            fill(sh, n.get('fill', '0B1A33'), n.get('alpha', 75))
            line(sh, n.get('border', TH['accent2']), n.get('border_w', 1.75))
            pad = n.get('pad', 0.015)
            text(sl, [x + pad, y + pad, w - 2 * pad, h - 2 * pad], n.get('lines', []), 16, 'text', 'body', False,
                 n.get('align', 'l'), n.get('anchor', 't'), spacing=4)

    def pt(p):
        if isinstance(p, list):
            return p
        i, side = p.split('.'); x, y, w, h = nodes[i]
        return {'l': [x, y + h / 2], 'r': [x + w, y + h / 2], 't': [x + w / 2, y], 'b': [x + w / 2, y + h]}[side]

    for L in d.get('links', []):
        pts = [pt(L['from'])] + L.get('via', []) + [pt(L['to'])]
        arr = L.get('arrows', 'both')
        for k in range(len(pts) - 1):
            (a, b) = pts[k], pts[k + 1]
            c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(a[0]), Y(a[1]), X(b[0]), Y(b[1]))
            line(c, L.get('color', '93C1E4'), L.get('width', 2), 'dash' if L.get('dash') else None)
            ln = c.line._get_or_add_ln()
            if k == 0 and arr in ('both', 'start'):
                etree.SubElement(ln, qn('a:headEnd')).set('type', 'triangle')
            if k == len(pts) - 2 and arr in ('both', 'end'):
                etree.SubElement(ln, qn('a:tailEnd')).set('type', 'triangle')
    for t in d.get('labels', []):
        text(sl, t['box'], t['text'], t.get('size', 15), t.get('color', 'muted'), t.get('font', 'body'), t.get('bold', False), t.get('align', 'l'))


def gfill(fmt, c1, c2, angle=90):
    f = fmt.fill; f.gradient(); f.gradient_angle = angle
    st = f.gradient_stops; st[0].position, st[1].position = 0.0, 1.0
    st[0].color.rgb, st[1].color.rgb = RGBColor.from_string(c1), RGBColor.from_string(c2)


def glow(fmt, color, rad=14, alpha=55):
    sp = fmt._element.get_or_add_spPr()
    for e in sp.findall(qn('a:effectLst')):
        sp.remove(e)
    g = etree.SubElement(etree.SubElement(sp, qn('a:effectLst')), qn('a:glow')); g.set('rad', str(rad * 12700))
    c = etree.SubElement(g, qn('a:srgbClr')); c.set('val', color); etree.SubElement(c, qn('a:alpha')).set('val', str(alpha * 1000))


def style_chart(sl, sh, c):
    """Alive charts: gradient bars/slices, glow on the point that proves the point, clean axes, white labels, donut center label."""
    ch = sh.chart; pl = ch.plots[0]; ser = pl.series[0]; n = len(list(ser.values)); hl = c.get('highlight')
    a1, a2 = TH['accent2'], TH['accent']
    dim = c.get('dim', [('2D6BB0', '1F4F8A'), ('3B7FC4', '2D6BB0'), ('5E8FBF', '3B7FC4'), ('93C1E4', '5E8FBF'), ('5E7FA6', '2C4466')])
    donut = any(k in str(ch.chart_type) for k in ('DOUGHNUT', 'PIE'))
    for i in range(n):
        f = ser.points[i].format
        if hl is None or i == hl:
            gfill(f, a1, a2, 45 if donut else 90)
            if i == hl:
                glow(f, a1, 16, 60)
        else:
            gfill(f, *dim[(i - 1) % len(dim)], 45 if donut else 0)
        if donut:
            f.line.color.rgb = RGBColor.from_string(c.get('gap', '07111F')); f.line.width = Pt(2)
        else:
            f.line.fill.background()
    if ch.has_title:
        for p in ch.chart_title.text_frame.paragraphs:
            for r in p.runs:
                r.font.name, r.font.size, r.font.bold = TH['display'], Pt(16), True; r.font.color.rgb = RGBColor(255, 255, 255)
    if pl.has_data_labels:
        dl = pl.data_labels; dl.font.size, dl.font.bold = Pt(14), True; dl.font.color.rgb = RGBColor(255, 255, 255)
    if ch.has_legend:
        ch.legend.font.size = Pt(13); ch.legend.font.color.rgb = RGBColor(0xE6, 0xED, 0xF5)
    if not donut:
        pl.gap_width = c.get('gap_width', 50)
        if c.get('clean_axes', True):
            ch.value_axis.has_major_gridlines = False; ch.value_axis.visible = False
    if c.get('center'):
        x0, y0, w0, h0 = sh.left / W, sh.top / H, sh.width / W, sh.height / H; cy = y0 + h0 * c.get('center_y', 0.36)
        text(sl, [x0 + w0 * 0.3, cy - 0.045, w0 * 0.4, 0.07], c['center'][0], 40, a1, 'display', True, 'c')
        if len(c['center']) > 1:
            text(sl, [x0 + w0 * 0.3, cy + 0.03, w0 * 0.4, 0.035], c['center'][1], 14, 'E6EDF5', 'body', False, 'c')


for op in P.get('remove', []):
    sl = S(op['slide'])
    for nm in op['shapes']:
        sh = next((s for s in sl.shapes if s.name == nm), None)
        if sh is None:
            print(f'warning: remove: no shape {nm!r} on slide {op["slide"]}', file=sys.stderr)
            continue
        sh._element.getparent().remove(sh._element)
for op in P.get('move', []):
    sh = find(S(op['slide']), op['shape']); b = op['box']
    sh.left, sh.top = X(b[0]), Y(b[1])
    if len(b) > 2 and b[2]:
        sh.width = X(b[2])
    if len(b) > 3 and b[3]:
        sh.height = Y(b[3])
for op in P.get('title_lines', []):
    title_lines(find(S(op['slide']), op['shape']), op['lines'], op['size'], op.get('box'), op.get('spacing', 95))
for op in P.get('center', []):
    center(find(S(op['slide']), op['shape']), op['box'], op.get('size'))
for op in P.get('align', []):
    for p in find(S(op['slide']), op['shape']).text_frame.paragraphs:
        p.alignment = AL[op.get('align', 'c')]
for op in P.get('logos', []):
    chip(S(op['slide']), op['box'], op['files'], op.get('name', 'logos'))
for op in P.get('cards', []):
    cards(S(op['slide']), op)
for op in P.get('diagram', []):
    diagram(S(op['slide']), op)
for op in P.get('text', []):
    text(S(op['slide']), op['box'], op['paras'], op.get('size', 18), op.get('color'), op.get('font'), op.get('bold', False),
         op.get('align', 'l'), op.get('anchor', 'm'))
for op in P.get('charts', []):
    sl = S(op['slide']); style_chart(sl, find(sl, op['name']), op)
if P.get('kpi', True):            # KPI tiles: centered text, one value size per row so no value wraps
    for sl in prs.slides:
        rows = {}
        for sh in sl.shapes:
            if not sh.has_text_frame or not re.search(r'-(value|label)-\d+$', sh.name):
                continue
            for p in sh.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
            if '-value-' in sh.name:
                fit = sh.width / 12700 * 0.95 / (max(1, len(sh.text_frame.text)) * 0.6)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        rows.setdefault(re.sub(r'-value-\d+$', '', sh.name), []).append((r, min(r.font.size.pt if r.font.size else 40, fit)))
        for runs in rows.values():
            m = min(v for _, v in runs)
            for r, _ in runs:
                r.font.size = Pt(m)
N = P.get('numbers', {})         # slide numbers: on by default, counting starts at the table of contents
if N is not False:
    first, n0, skip = N.get('start_slide', 2), N.get('start_at', 1), set(N.get('skip', []))
    bx = N.get('box', [0.925, 0.925, 0.055, 0.045])
    total = len(prs.slides) - first + n0

    def own_number(s):              # the template's own page number: reuse it instead of adding a second one
        if s.name == 'slide-number-tpl':
            return True
        if s.is_placeholder and s.placeholder_format.type is not None and int(s.placeholder_format.type) == 13:
            return True
        if not s.has_text_frame or s.name.startswith(('card-', 'kpi', 'node-', 'logo', 'chart', 'slide-number')):
            return False
        if s._element.findall('.//' + qn('a:fld')) and any(x.get('type') == 'slidenum' for x in s._element.iter(qn('a:fld'))):
            return True
        tx = s.text_frame.text.strip()
        if not re.fullmatch(r'\d{1,3}', tx) or s.width * s.height > 0.01 * W * H:
            return False
        cx, cy = (s.left + s.width / 2) / W, (s.top + s.height / 2) / H
        return cy > 0.82 or cy < 0.12 or cx > 0.92 or cx < 0.08

    def set_text(s, value):
        for fld in list(s._element.iter(qn('a:fld'))):      # a slide-number field would ignore our numbering
            r = etree.Element(qn('a:r'))
            for c in list(fld):
                if c.tag in (qn('a:rPr'), qn('a:t')):
                    r.append(c)
            if r.find(qn('a:t')) is None:
                etree.SubElement(r, qn('a:t'))
            fld.getparent().replace(fld, r)
        runs = [r for p in s.text_frame.paragraphs for r in p.runs]
        if runs:
            runs[0].text = value
            for r in runs[1:]:
                r.text = ''
        else:
            s.text_frame.text = value
        s.name = 'slide-number-tpl'

    found = {i: [s for s in sl.shapes if own_number(s)] for i, sl in enumerate(prs.slides, 1)}
    allf = [s for ss in found.values() for s in ss]
    proto = next((s for s in allf if s.text_frame.text.strip()), allf[0] if allf else None)
    pad = 0
    if proto is not None:
        t0 = proto.text_frame.text.strip()
        pad = len(t0) if t0.startswith('0') and len(t0) > 1 else int(N.get('pad', 0))
    for i, sl in enumerate(prs.slides, 1):
        for sh in [s for s in sl.shapes if s.name == 'slide-number']:
            sh._element.getparent().remove(sh._element)
        numbered = i >= first and i not in skip
        value = str(i - first + n0).zfill(pad) if numbered else ''
        if proto is not None:                     # template style: rewrite its numbers, copy one where missing
            if not found[i] and numbered:
                el = copy.deepcopy(proto._element)
                sl.shapes._spTree.append(el)
                found[i] = [sl.shapes[-1]]
            for s in found[i]:
                set_text(s, value)
            continue
        if not numbered:
            continue
        b = list(bx)
        for s in sl.shapes:          # step above a logo chip or added content sitting in the corner
            if s.name.startswith(('logos', 'logo-', 'node-', 'card-')):
                sx, sy, sw, shh = s.left / W, s.top / H, s.width / W, s.height / H
                if sx < b[0] + b[2] and sx + sw > b[0] and sy < b[1] + b[3] and sy + shh > b[1]:
                    b[1] = sy - b[3] - 0.01
        text(sl, b, N.get('format', '{n}').format(n=i - first + n0, total=total),
             N.get('size', 18), N.get('color', 'text'), N.get('font', 'display'), N.get('bold', True), 'r', 'm', name='slide-number')
C = P.get('canvas')               # banners: crop the 16:9 slide to a wide strip after the layout is placed
if C:
    aspect = float(C.get('aspect', 3))
    if aspect <= 1:
        sys.exit('canvas: a banner is always wider than tall (aspect > 1)')
    newW, newH = (W, int(W / aspect)) if aspect >= W / H else (int(H * aspect), H)
    dx, dy = int(C.get('left', 0) * W), int(C.get('top', 0) * H)
    if dy + newH > H + 1 or dx + newW > W + 1:
        sys.exit('canvas: the crop band runs past the slide (lower "top" or "left")')
    for sl in prs.slides:
        for sh in sl.shapes:
            sh.left, sh.top = sh.left - dx, sh.top - dy
    prs.slide_width, prs.slide_height = newW, newH
prs.save(out)
print(json.dumps({'polished': out, 'ops': {k: len(v) for k, v in P.items() if isinstance(v, list)}}))
